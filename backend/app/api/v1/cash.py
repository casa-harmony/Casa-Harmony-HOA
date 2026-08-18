from __future__ import annotations

import csv
import io
import uuid
from datetime import date

from fastapi import APIRouter, Depends, Form, HTTPException, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import Principal, require_active_tenant, require_permission
from app.models.cash import CeBankAccount, CeStatementHeader
from app.models.identity import Tenant
from app.schemas.cash import (
    AdjustIn,
    BankAccountCreate,
    BankAccountOut,
    CashPositionOut,
    MatchIn,
    StatementCreate,
    StatementDetail,
    StatementOut,
)
from app.services import audit, cash_mgmt, cash_reports
from app.services.cash_mgmt import CashError

router = APIRouter(prefix="/cash", tags=["cash"], dependencies=[Depends(require_active_tenant)])
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _err(exc: CashError):
    raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc))


def _bank_account_out(acct: CeBankAccount) -> BankAccountOut:
    masked = f"****{acct.account_number[-4:]}" if acct.account_number and len(acct.account_number) >= 4 else None
    return BankAccountOut(
        id=acct.id, account_code=acct.account_code, name=acct.name, fund_value=acct.fund_value,
        bank_name=acct.bank_name, routing_number=acct.routing_number,
        gl_cash_combination_id=acct.gl_cash_combination_id, currency=acct.currency, active=acct.active,
        bank=acct.bank_name, fund=acct.fund_value, masked=masked,
    )


# --- Bank accounts ---------------------------------------------------------
@router.get("/bank-accounts", response_model=list[BankAccountOut])
def list_accounts(db: Session = Depends(get_db),
                  p: Principal = Depends(require_permission("cash.manage"))):
    rows = db.execute(select(CeBankAccount).where(CeBankAccount.tenant_id == p.tenant_id)
                      .order_by(CeBankAccount.fund_value, CeBankAccount.account_code)).scalars().all()
    return [_bank_account_out(a) for a in rows]


@router.post("/bank-accounts", response_model=BankAccountOut, status_code=status.HTTP_201_CREATED)
def create_account(payload: BankAccountCreate, db: Session = Depends(get_db),
                   p: Principal = Depends(require_permission("cash.manage"))):
    try:
        acct = cash_mgmt.create_bank_account(db, tenant_id=p.tenant_id, created_by=p.user.id,
                                             **payload.model_dump())
    except CashError as exc:
        _err(exc)
    audit.record(db, action="CREATE", entity_type="CeBankAccount", entity_id=acct.id,
                 after={"account_code": acct.account_code})
    return _bank_account_out(acct)


# --- Statements ------------------------------------------------------------
@router.get("/statements", response_model=list[StatementOut])
def list_statements(ce_bank_account_id: uuid.UUID | None = None, db: Session = Depends(get_db),
                    p: Principal = Depends(require_permission("cash.manage"))):
    stmt = select(CeStatementHeader).where(CeStatementHeader.tenant_id == p.tenant_id)
    if ce_bank_account_id:
        stmt = stmt.where(CeStatementHeader.ce_bank_account_id == ce_bank_account_id)
    return db.execute(stmt.order_by(CeStatementHeader.statement_date.desc())).scalars().all()


@router.get("/statements/{statement_id}", response_model=StatementDetail)
def get_statement(statement_id: uuid.UUID, db: Session = Depends(get_db),
                  p: Principal = Depends(require_permission("cash.manage"))):
    h = db.get(CeStatementHeader, statement_id)
    if h is None or h.tenant_id != p.tenant_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Statement not found")
    return h


@router.post("/statements", response_model=StatementDetail, status_code=status.HTTP_201_CREATED)
def create_statement(payload: StatementCreate, db: Session = Depends(get_db),
                     p: Principal = Depends(require_permission("cash.manage"))):
    try:
        h = cash_mgmt.import_statement(
            db, tenant_id=p.tenant_id, ce_bank_account_id=payload.ce_bank_account_id,
            statement_date=payload.statement_date, opening_balance=payload.opening_balance,
            closing_balance=payload.closing_balance,
            lines=[ln.model_dump() for ln in payload.lines], created_by=p.user.id)
    except CashError as exc:
        _err(exc)
    audit.record(db, action="CREATE", entity_type="CeStatementHeader", entity_id=h.id)
    return h


def _parse_amount(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").replace("$", ""))
    except ValueError:
        return None


@router.post("/statements/import", response_model=StatementDetail, status_code=status.HTTP_201_CREATED)
async def import_statement_file(
    file: UploadFile,
    ce_bank_account_id: uuid.UUID = Form(...),
    statement_date: date = Form(...),
    opening_balance: float = Form(0),
    closing_balance: float = Form(0),
    db: Session = Depends(get_db),
    p: Principal = Depends(require_permission("cash.manage")),
):
    """Import a bank statement from CSV or xlsx. Columns: date, description, reference, amount."""
    raw = await file.read()
    records: list[dict] = []
    if file.filename and file.filename.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(raw), data_only=True).active
        headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            records.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        for rec in csv.DictReader(io.StringIO(text)):
            records.append({(k or "").strip().lower(): v for k, v in rec.items()})

    lines = []
    for rec in records:
        amt = _parse_amount(rec.get("amount"))
        if amt is None:
            continue
        ld = rec.get("date")
        try:
            ld = date.fromisoformat(str(ld)[:10]) if ld else None
        except ValueError:
            ld = None
        lines.append({"line_date": ld, "description": rec.get("description"),
                      "reference": rec.get("reference"), "amount": amt})
    if not lines:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "No valid statement lines found")
    try:
        h = cash_mgmt.import_statement(
            db, tenant_id=p.tenant_id, ce_bank_account_id=ce_bank_account_id,
            statement_date=statement_date, opening_balance=opening_balance,
            closing_balance=closing_balance, lines=lines, created_by=p.user.id)
    except CashError as exc:
        _err(exc)
    audit.record(db, action="IMPORT", entity_type="CeStatementHeader", entity_id=h.id,
                 after={"lines": len(lines)})
    return h


# --- Reconciliation --------------------------------------------------------
@router.post("/statements/lines/{line_id}/match", response_model=StatementDetail)
def match(line_id: uuid.UUID, payload: MatchIn, db: Session = Depends(get_db),
          p: Principal = Depends(require_permission("cash.manage"))):
    try:
        ln = cash_mgmt.match_line(db, tenant_id=p.tenant_id, line_id=line_id,
                                  payment_id=payload.payment_id, receipt_id=payload.receipt_id)
    except CashError as exc:
        _err(exc)
    audit.record(db, action="MATCH", entity_type="CeStatementLine", entity_id=ln.id)
    return db.get(CeStatementHeader, ln.header_id)


@router.post("/statements/lines/{line_id}/adjust", response_model=StatementDetail)
def adjust(line_id: uuid.UUID, payload: AdjustIn, db: Session = Depends(get_db),
           p: Principal = Depends(require_permission("cash.manage"))):
    try:
        ln = cash_mgmt.create_adjustment(
            db, tenant_id=p.tenant_id, line_id=line_id,
            offset_combination_id=payload.offset_combination_id, gl_date=payload.gl_date,
            description=payload.description, created_by=p.user.id)
    except CashError as exc:
        _err(exc)
    audit.record(db, action="ADJUST", entity_type="CeStatementLine", entity_id=ln.id)
    return db.get(CeStatementHeader, ln.header_id)


@router.post("/statements/{statement_id}/reconcile", response_model=StatementOut)
def reconcile(statement_id: uuid.UUID, db: Session = Depends(get_db),
              p: Principal = Depends(require_permission("cash.manage"))):
    h = db.get(CeStatementHeader, statement_id)
    if h is None or h.tenant_id != p.tenant_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Statement not found")
    try:
        cash_mgmt.reconcile_statement(db, h, p.user.id)
    except CashError as exc:
        _err(exc)
    audit.record(db, action="RECONCILE", entity_type="CeStatementHeader", entity_id=h.id)
    return h


# --- Dashboard + reports ---------------------------------------------------
@router.get("/position", response_model=list[CashPositionOut])
def position(db: Session = Depends(get_db),
             p: Principal = Depends(require_permission("cash.manage"))):
    return cash_mgmt.cash_position(db, p.tenant_id)


def _tname(db, tid):
    t = db.get(Tenant, tid)
    return t.name if t else "HOA"


@router.get("/reports/bank-rec/export")
def export_bank_rec(ce_bank_account_id: uuid.UUID | None = None, db: Session = Depends(get_db),
                    p: Principal = Depends(require_permission("report.read"))):
    content = cash_reports.build_bank_rec_report(db, p.tenant_id, _tname(db, p.tenant_id), ce_bank_account_id)
    return Response(content=content, media_type=XLSX,
                    headers={"Content-Disposition": 'attachment; filename="bank_reconciliation.xlsx"'})


@router.get("/reports/cash-flow/export")
def export_cash_flow(start: date, end: date, db: Session = Depends(get_db),
                     p: Principal = Depends(require_permission("report.read"))):
    content = cash_reports.build_cash_flow_summary(db, p.tenant_id, start, end, _tname(db, p.tenant_id))
    return Response(content=content, media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="cash_flow_{start}_{end}.xlsx"'})
