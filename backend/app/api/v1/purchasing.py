from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile, status
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import Principal, require_active_tenant, require_permission
from app.models.identity import Tenant
from app.models.kff import GlCodeCombination
from app.models.masters import ApSupplier
from app.models.procurement import PoDistribution, PoHeader, PoLine
from app.models.workflow import ApprovalRequest
from app.schemas.approvals import ActIn
from app.schemas.procurement import PoCreate, PoDetailOut, PoOut
from app.services import approvals, audit, budgeting, encumbrance
from app.services.budgeting import BudgetError
from app.services.distributions import DistributionError
from app.services.po_service import create_po

router = APIRouter(
    prefix="/purchasing", tags=["purchasing"], dependencies=[Depends(require_active_tenant)]
)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _finalize(po: PoHeader, approver_id: uuid.UUID) -> None:
    po.status = "APPROVED"
    po.approval_status = "APPROVED"
    po.approved_by = approver_id
    po.approved_at = datetime.now(timezone.utc)


def _get_po(db: Session, po_id: uuid.UUID, tenant_id: uuid.UUID) -> PoHeader:
    po = db.get(PoHeader, po_id)
    if po is None or po.tenant_id != tenant_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "PO not found")
    return po


@router.get("", response_model=list[PoOut])
def list_pos(
    vendor_id: uuid.UUID | None = None,
    cost_center: str | None = None,
    po_status: str | None = None,
    active_on: date | None = None,
    min_remaining: float | None = None,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.manage")),
):
    stmt = select(PoHeader).where(PoHeader.tenant_id == principal.tenant_id)
    if vendor_id:
        stmt = stmt.where(PoHeader.vendor_id == vendor_id)
    if po_status:
        stmt = stmt.where(PoHeader.status == po_status)
    if active_on:
        stmt = stmt.where(
            ((PoHeader.start_date.is_(None)) | (PoHeader.start_date <= active_on)),
            ((PoHeader.end_date.is_(None)) | (PoHeader.end_date >= active_on)),
        )
    if min_remaining is not None:
        stmt = stmt.where((PoHeader.amount_limit - PoHeader.billed_amount) >= min_remaining)
    if cost_center:
        # POs that have any distribution on the given cost center.
        sub = (
            select(PoDistribution.id)
            .join(PoLine, PoLine.id == PoDistribution.po_line_id)
            .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id)
            .where(PoLine.po_header_id == PoHeader.id,
                   GlCodeCombination.cost_center_value == cost_center)
        )
        stmt = stmt.where(sub.exists())
    return db.execute(stmt.order_by(PoHeader.po_number.desc())).scalars().all()


@router.get("/{po_id}", response_model=PoDetailOut)
def get_po(
    po_id: uuid.UUID,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.manage")),
):
    return _get_po(db, po_id, principal.tenant_id)


@router.post("/{po_id}/close", response_model=PoOut)
def close_po(
    po_id: uuid.UUID,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.approve")),
):
    po = _get_po(db, po_id, principal.tenant_id)
    if po.status in ("INCOMPLETE", "CANCELLED"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Cannot close a {po.status} PO")
    po.status = "CLOSED"
    audit.record(db, action="CLOSE", entity_type="PoHeader", entity_id=po.id)
    return po


@router.post("", response_model=PoOut, status_code=status.HTTP_201_CREATED)
def create(
    payload: PoCreate,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.manage")),
):
    try:
        po = create_po(
            db, tenant_id=principal.tenant_id, vendor_id=payload.vendor_id,
            order_date=payload.order_date, description=payload.description,
            document_type=payload.document_type, start_date=payload.start_date,
            end_date=payload.end_date, amount_limit=payload.amount_limit,
            lines=[ln.model_dump() for ln in payload.lines],
            created_by=principal.user.id,
        )
    except DistributionError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc))
    audit.record(db, action="CREATE", entity_type="PoHeader", entity_id=po.id,
                 after={"po_number": po.po_number, "amount": str(po.amount)})
    return po


@router.post("/{po_id}/submit", response_model=PoOut)
def submit(
    po_id: uuid.UUID,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.manage")),
):
    po = _get_po(db, po_id, principal.tenant_id)
    if po.status not in ("INCOMPLETE", "REJECTED"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Cannot submit a {po.status} PO")
    req = approvals.submit(
        db, tenant_id=principal.tenant_id, document_type="PO", document_id=po.id,
        amount=po.amount, submitted_by=principal.user.id,
    )
    po.status = "SUBMITTED"
    po.approval_status = "PENDING"
    if req.status == "APPROVED":  # no hierarchy → auto-approved
        _finalize(po, principal.user.id)
        encumbrance.encumber_po(db, po, principal.user.id)
        try:
            budgeting.enforce_po_budget(db, po, principal.user.id)
        except BudgetError as exc:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc))
    audit.record(db, action="SUBMIT", entity_type="PoHeader", entity_id=po.id)
    return po


@router.post("/{po_id}/approve", response_model=PoOut)
def approve(
    po_id: uuid.UUID,
    payload: ActIn,
    db: Session = Depends(get_db),
    principal: Principal = Depends(require_permission("po.approve")),
):
    po = _get_po(db, po_id, principal.tenant_id)
    req = db.execute(
        select(ApprovalRequest).where(
            ApprovalRequest.tenant_id == principal.tenant_id,
            ApprovalRequest.document_type == "PO", ApprovalRequest.document_id == po.id,
        )
    ).scalar_one_or_none()
    if req is None or req.status != "PENDING":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No pending approval for this PO")
    approvals.act(db, request=req, approver_id=principal.user.id,
                  approve=payload.approve, comments=payload.comments)
    if req.status == "APPROVED":
        _finalize(po, principal.user.id)
        encumbrance.encumber_po(db, po, principal.user.id)
        try:
            budgeting.enforce_po_budget(db, po, principal.user.id)
        except BudgetError as exc:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc))
    elif req.status == "REJECTED":
        po.status = "REJECTED"
        po.approval_status = "REJECTED"
    audit.record(db, action="APPROVE" if payload.approve else "REJECT",
                 entity_type="PoHeader", entity_id=po.id)
    return po


# --- xlsx export / import --------------------------------------------------
EXPORT_HEADERS = [
    "po_number", "document_type", "vendor_number", "vendor_name", "order_date",
    "start_date", "end_date", "amount", "amount_limit", "billed_amount", "status",
    "line_num", "item_description", "quantity", "unit_price", "account", "dist_amount", "fund",
]


@router.get("/export/xlsx")
def export_pos(db: Session = Depends(get_db),
               principal: Principal = Depends(require_permission("po.manage"))):
    rows = db.execute(
        select(PoHeader, PoLine, PoDistribution, GlCodeCombination, ApSupplier)
        .join(ApSupplier, ApSupplier.id == PoHeader.vendor_id)
        .join(PoLine, PoLine.po_header_id == PoHeader.id, isouter=True)
        .join(PoDistribution, PoDistribution.po_line_id == PoLine.id, isouter=True)
        .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id, isouter=True)
        .where(PoHeader.tenant_id == principal.tenant_id)
        .order_by(PoHeader.po_number, PoLine.line_num, PoDistribution.distribution_num)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Purchase Orders"
    ws.append(EXPORT_HEADERS)
    for po, line, dist, cc, v in rows:
        ws.append([
            po.po_number, po.document_type, v.vendor_number, v.name,
            po.order_date.isoformat() if po.order_date else "",
            po.start_date.isoformat() if po.start_date else "",
            po.end_date.isoformat() if po.end_date else "",
            float(po.amount), float(po.amount_limit), float(po.billed_amount), po.status,
            line.line_num if line else "", line.item_description if line else "",
            float(line.quantity) if line else "", float(line.unit_price) if line else "",
            cc.concatenated_segments if cc else "", float(dist.amount) if dist else "",
            dist.fund_value if dist else "",
        ])
    buf = io.BytesIO(); wb.save(buf)
    return Response(content=buf.getvalue(), media_type=XLSX,
                    headers={"Content-Disposition": 'attachment; filename="purchase_orders.xlsx"'})


@router.post("/import/xlsx")
async def import_pos(file: UploadFile,
                     db: Session = Depends(get_db),
                     principal: Principal = Depends(require_permission("po.manage"))):
    """Import POs. Each row is one line+distribution; rows are grouped by the 'ref'
    column into a single PO. Columns: ref, vendor_number, document_type, order_date,
    start_date, end_date, amount_limit, item_description, quantity, unit_price, account."""
    wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]

    def col(row, name):
        return row[headers.index(name)] if name in headers and headers.index(name) < len(row) else None

    groups: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        ref = str(col(row, "ref") or "")
        if not ref:
            continue
        acct = str(col(row, "account") or "").strip()
        cc = db.execute(select(GlCodeCombination).where(
            GlCodeCombination.tenant_id == principal.tenant_id,
            GlCodeCombination.concatenated_segments == acct)).scalar_one_or_none()
        if cc is None:
            continue
        line_amt = (float(col(row, "quantity") or 1) * float(col(row, "unit_price") or 0))
        line = {"item_description": str(col(row, "item_description") or "Item"),
                "quantity": col(row, "quantity") or 1, "unit_price": col(row, "unit_price") or 0,
                "distributions": [{"code_combination_id": cc.id, "amount": round(line_amt, 2)}]}
        if ref not in groups:
            vno = str(col(row, "vendor_number") or "")
            vendor = db.execute(select(ApSupplier).where(
                ApSupplier.tenant_id == principal.tenant_id,
                ApSupplier.vendor_number == vno)).scalar_one_or_none()
            groups[ref] = {"vendor": vendor, "row": row, "lines": []}
        groups[ref]["lines"].append(line)

    created, skipped = 0, 0
    for ref, g in groups.items():
        if g["vendor"] is None:
            skipped += 1
            continue
        row = g["row"]
        def d(name):
            v = col(row, name)
            return date.fromisoformat(str(v)[:10]) if v else None
        try:
            create_po(db, tenant_id=principal.tenant_id, vendor_id=g["vendor"].id,
                      order_date=d("order_date") or date.today(),
                      document_type=str(col(row, "document_type") or "STANDARD").upper(),
                      start_date=d("start_date"), end_date=d("end_date"),
                      amount_limit=col(row, "amount_limit"),
                      lines=g["lines"], created_by=principal.user.id)
            created += 1
        except DistributionError:
            skipped += 1
    audit.record(db, action="IMPORT", entity_type="PoHeader",
                 after={"created": created, "skipped": skipped})
    return {"created": created, "skipped": skipped}


# --- Reports ---------------------------------------------------------------
from app.services import po_reports  # noqa: E402


def _tname(db, tid):
    t = db.get(Tenant, tid)
    return t.name if t else "HOA"


def _xlsx(content: bytes, filename: str) -> Response:
    return Response(content=content, media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/run-budget-checks")
def run_budget_checks_now(db: Session = Depends(get_db),
                          p: Principal = Depends(require_permission("po.manage"))):
    """Manually trigger the PO budget-check sweep for this tenant."""
    from app.services.matching import run_budget_checks
    raised = run_budget_checks(db, p.tenant_id)
    return {"alerts_raised": raised}


@router.get("/reports/contract-utilization")
def report_contract_util(db: Session = Depends(get_db),
                         p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_contract_utilization(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "contract_utilization.xlsx")


@router.get("/reports/commitment-register")
def report_commitment(db: Session = Depends(get_db),
                      p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_commitment_register(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "po_commitment_register.xlsx")


@router.get("/reports/variance-by-contract")
def report_variance(db: Session = Depends(get_db),
                    p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_variance_by_contract(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "contract_variance.xlsx")


@router.get("/reports/cost-center-utilization")
def report_cc_util(db: Session = Depends(get_db),
                   p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_cost_center_utilization(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "cost_center_utilization.xlsx")


@router.get("/reports/matched-summary")
def report_matched(db: Session = Depends(get_db),
                   p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_matched_summary(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "matched_vs_nonmatched.xlsx")


@router.get("/reports/activity-log")
def report_activity(db: Session = Depends(get_db),
                    p: Principal = Depends(require_permission("report.read"))):
    return _xlsx(po_reports.build_po_activity_log(db, p.tenant_id, _tname(db, p.tenant_id)),
                 "po_activity_log.xlsx")
