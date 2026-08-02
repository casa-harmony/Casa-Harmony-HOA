"""Cash Management reports: bank reconciliation report + cash flow summary by Fund."""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from app.models.cash import CeBankAccount, CeStatementHeader, CeStatementLine
from app.services.reports import _BOLD, _header_row


def build_bank_rec_report(db, tenant_id, tenant_name, ce_bank_account_id=None) -> bytes:
    stmt = (
        select(CeStatementHeader, CeBankAccount)
        .join(CeBankAccount, CeBankAccount.id == CeStatementHeader.ce_bank_account_id)
        .where(CeStatementHeader.tenant_id == tenant_id)
    )
    if ce_bank_account_id:
        stmt = stmt.where(CeStatementHeader.ce_bank_account_id == ce_bank_account_id)
    headers = db.execute(stmt.order_by(CeStatementHeader.statement_date)).all()

    wb = Workbook(); ws = wb.active; ws.title = "Bank Reconciliation"
    ws["A1"] = f"{tenant_name} — Bank Reconciliation Report"
    ws["A1"].font = Font(size=14, bold=True)
    r = 3
    for hdr, acct in headers:
        ws.cell(row=r, column=1, value=f"{acct.account_code} — {acct.name} (Fund {acct.fund_value})").font = _BOLD
        r += 1
        ws.cell(row=r, column=1, value=f"Statement {hdr.statement_date.isoformat()} · status {hdr.status}")
        r += 1
        ws.cell(row=r, column=1, value=f"Opening {hdr.opening_balance}  Closing {hdr.closing_balance}")
        r += 1
        _header_row(ws, r, ["Date", "Description", "Reference", "Amount", "Reconciled", "Via"])
        r += 1
        rec = unrec = Decimal("0")
        for ln in hdr.lines:
            ws.cell(row=r, column=1, value=ln.line_date.isoformat() if ln.line_date else "")
            ws.cell(row=r, column=2, value=ln.description or "")
            ws.cell(row=r, column=3, value=ln.reference or "")
            ws.cell(row=r, column=4, value=float(ln.amount))
            ws.cell(row=r, column=5, value="Y" if ln.reconciled else "N")
            ws.cell(row=r, column=6, value=ln.match_type or "")
            (rec := rec + Decimal(ln.amount)) if ln.reconciled else (unrec := unrec + Decimal(ln.amount))
            r += 1
        ws.cell(row=r, column=2, value="Reconciled / Unreconciled").font = _BOLD
        ws.cell(row=r, column=4, value=f"{rec} / {unrec}").font = _BOLD
        r += 2
    for col, w in {"A": 14, "B": 32, "C": 16, "D": 14, "E": 11, "F": 12}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_cash_flow_summary(db, tenant_id, start: date, end: date, tenant_name) -> bytes:
    """Inflows/outflows by Fund from statement lines in a date range."""
    rows = db.execute(
        select(CeStatementLine, CeBankAccount)
        .join(CeStatementHeader, CeStatementHeader.id == CeStatementLine.header_id)
        .join(CeBankAccount, CeBankAccount.id == CeStatementHeader.ce_bank_account_id)
        .where(CeStatementLine.tenant_id == tenant_id,
               CeStatementHeader.statement_date >= start,
               CeStatementHeader.statement_date <= end)
    ).all()
    inflow: dict = defaultdict(Decimal)
    outflow: dict = defaultdict(Decimal)
    for ln, acct in rows:
        amt = Decimal(ln.amount)
        if amt >= 0:
            inflow[acct.fund_value] += amt
        else:
            outflow[acct.fund_value] += amt

    wb = Workbook(); ws = wb.active; ws.title = "Cash Flow by Fund"
    ws["A1"] = f"{tenant_name} — Cash Flow Summary {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Fund", "Inflows", "Outflows", "Net"])
    r = 4
    funds = sorted(set(inflow) | set(outflow))
    ti = to = Decimal("0")
    for f in funds:
        i, o = inflow.get(f, Decimal("0")), outflow.get(f, Decimal("0"))
        ws.cell(row=r, column=1, value=f)
        ws.cell(row=r, column=2, value=float(i))
        ws.cell(row=r, column=3, value=float(o))
        ws.cell(row=r, column=4, value=float(i + o))
        ti += i; to += o
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = _BOLD
    ws.cell(row=r, column=2, value=float(ti)).font = _BOLD
    ws.cell(row=r, column=3, value=float(to)).font = _BOLD
    ws.cell(row=r, column=4, value=float(ti + to)).font = _BOLD
    for col, w in {"A": 12, "B": 16, "C": 16, "D": 16}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
