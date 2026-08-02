"""AP payment reports: payment register, aged payables, cash requirements (by Fund)."""
from __future__ import annotations

import io
import uuid
from collections import defaultdict
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.kff import GlCodeCombination
from app.models.masters import ApSupplier
from app.models.payables import ApInvoice, ApInvoiceDistribution, ApInvoiceLine
from app.models.payments import ApPayment, ApPaymentSchedule
from app.services.periods import period_name
from app.services.reports import _BOLD, _header_row

BUCKETS = ["Current", "1-30", "31-60", "61-90", "90+"]


def _bucket(due: date | None, as_of: date) -> str:
    if due is None or due >= as_of:
        return "Current"
    days = (as_of - due).days
    if days <= 30:
        return "1-30"
    if days <= 60:
        return "31-60"
    if days <= 90:
        return "61-90"
    return "90+"


def build_payment_register_workbook(db, tenant_id, start: date, end: date, tenant_name) -> bytes:
    rows = db.execute(
        select(ApPayment, ApSupplier)
        .join(ApSupplier, ApSupplier.id == ApPayment.vendor_id)
        .where(ApPayment.tenant_id == tenant_id,
               ApPayment.payment_date >= start, ApPayment.payment_date <= end)
        .order_by(ApPayment.payment_date)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Payment Register"
    ws["A1"] = f"{tenant_name} — Payment Register {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Date", "Payment #", "Vendor", "Reference", "Status", "Amount"])
    r = 4
    total = Decimal("0")
    for p, v in rows:
        ws.cell(row=r, column=1, value=p.payment_date.isoformat())
        ws.cell(row=r, column=2, value=p.payment_number)
        ws.cell(row=r, column=3, value=v.name)
        ws.cell(row=r, column=4, value=p.reference or "")
        ws.cell(row=r, column=5, value=p.status)
        ws.cell(row=r, column=6, value=float(p.amount))
        if p.status == "CREATED":
            total += Decimal(p.amount)
        r += 1
    ws.cell(row=r, column=5, value="TOTAL (active)").font = _BOLD
    ws.cell(row=r, column=6, value=float(total)).font = _BOLD
    for col, w in {"A": 12, "B": 16, "C": 28, "D": 16, "E": 10, "F": 14}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _open_with_invoice(db, tenant_id):
    return db.execute(
        select(ApPaymentSchedule, ApInvoice, ApSupplier)
        .join(ApInvoice, ApInvoice.id == ApPaymentSchedule.invoice_id)
        .join(ApSupplier, ApSupplier.id == ApInvoice.vendor_id)
        .where(ApPaymentSchedule.tenant_id == tenant_id, ApPaymentSchedule.status != "PAID")
    ).all()


def build_aged_payables_workbook(db, tenant_id, as_of: date, tenant_name) -> bytes:
    per_vendor: dict = {}
    totals = {b: Decimal("0") for b in BUCKETS}
    for sched, inv, v in _open_with_invoice(db, tenant_id):
        rem = sched.amount_remaining
        if rem <= 0:
            continue
        b = _bucket(sched.due_date, as_of)
        e = per_vendor.setdefault(v.id, {"name": v.name, "buckets": {k: Decimal("0") for k in BUCKETS}, "total": Decimal("0")})
        e["buckets"][b] += rem
        e["total"] += rem
        totals[b] += rem
    wb = Workbook(); ws = wb.active; ws.title = "Aged Payables"
    ws["A1"] = f"{tenant_name} — Aged Payables as of {as_of.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Vendor", *BUCKETS, "Total"])
    r = 4
    grand = Decimal("0")
    for e in sorted(per_vendor.values(), key=lambda x: x["name"]):
        ws.cell(row=r, column=1, value=e["name"])
        for i, b in enumerate(BUCKETS, start=2):
            ws.cell(row=r, column=i, value=float(e["buckets"][b]))
        ws.cell(row=r, column=7, value=float(e["total"]))
        grand += e["total"]
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = _BOLD
    for i, b in enumerate(BUCKETS, start=2):
        ws.cell(row=r, column=i, value=float(totals[b])).font = _BOLD
    ws.cell(row=r, column=7, value=float(grand)).font = _BOLD
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_cash_requirements_workbook(db, tenant_id, as_of: date, tenant_name) -> bytes:
    """Cash needed by Fund to clear open payables (remaining split by invoice fund)."""
    by_fund: dict[str, Decimal] = defaultdict(Decimal)
    for sched, inv, _v in _open_with_invoice(db, tenant_id):
        rem = sched.amount_remaining
        if rem <= 0 or not inv.amount:
            continue
        ratio = rem / Decimal(inv.amount)
        dists = db.execute(
            select(ApInvoiceDistribution)
            .join(ApInvoiceLine, ApInvoiceLine.id == ApInvoiceDistribution.invoice_line_id)
            .where(ApInvoiceLine.invoice_id == inv.id)
        ).scalars().all()
        for d in dists:
            by_fund[d.fund_value] += Decimal(d.amount) * ratio
    wb = Workbook(); ws = wb.active; ws.title = "Cash Requirements"
    ws["A1"] = f"{tenant_name} — Cash Requirements by Fund as of {as_of.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Fund", "Cash Required"])
    r = 4
    total = Decimal("0")
    for fund, amt in sorted(by_fund.items()):
        amt = amt.quantize(Decimal("0.01"))
        ws.cell(row=r, column=1, value=fund)
        ws.cell(row=r, column=2, value=float(amt))
        total += amt
        r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = _BOLD
    ws.cell(row=r, column=2, value=float(total)).font = _BOLD
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 18
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_invoice_register_workbook(db, tenant_id, start: date, end: date, tenant_name) -> bytes:
    """AP invoice register: invoices by invoice_date with vendor, tax, status, holds."""
    rows = db.execute(
        select(ApInvoice, ApSupplier)
        .join(ApSupplier, ApSupplier.id == ApInvoice.vendor_id)
        .where(ApInvoice.tenant_id == tenant_id,
               ApInvoice.invoice_date >= start, ApInvoice.invoice_date <= end)
        .order_by(ApInvoice.invoice_date)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Invoice Register"
    ws["A1"] = f"{tenant_name} — AP Invoice Register {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Date", "Invoice #", "Vendor", "Amount", "Tax", "Status", "Match", "Hold"])
    r = 4
    total = Decimal("0")
    for inv, v in rows:
        ws.cell(row=r, column=1, value=inv.invoice_date.isoformat())
        ws.cell(row=r, column=2, value=inv.invoice_number)
        ws.cell(row=r, column=3, value=v.name)
        ws.cell(row=r, column=4, value=float(inv.amount))
        ws.cell(row=r, column=5, value=float(inv.tax_amount or 0))
        ws.cell(row=r, column=6, value=inv.status)
        ws.cell(row=r, column=7, value=inv.match_status)
        ws.cell(row=r, column=8, value="HOLD" if inv.on_hold else "")
        if inv.status != "CANCELLED":
            total += Decimal(inv.amount)
        r += 1
    ws.cell(row=r, column=3, value="TOTAL").font = _BOLD
    ws.cell(row=r, column=4, value=float(total)).font = _BOLD
    for col, w in {"A": 12, "B": 16, "C": 28, "D": 14, "E": 10, "F": 12, "G": 16, "H": 8}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_ap_distributions_workbook(db, tenant_id, start: date, end: date, tenant_name) -> bytes:
    """AP distributions detail + a Fund × Period summary (GL date basis)."""
    rows = db.execute(
        select(ApInvoiceDistribution, ApInvoice, GlCodeCombination)
        .join(ApInvoiceLine, ApInvoiceLine.id == ApInvoiceDistribution.invoice_line_id)
        .join(ApInvoice, ApInvoice.id == ApInvoiceLine.invoice_id)
        .join(GlCodeCombination, GlCodeCombination.id == ApInvoiceDistribution.code_combination_id)
        .where(ApInvoice.tenant_id == tenant_id,
               ApInvoice.gl_date >= start, ApInvoice.gl_date <= end,
               ApInvoice.status != "CANCELLED")
        .order_by(ApInvoiceDistribution.fund_value)
    ).all()
    wb = Workbook()
    # Detail sheet
    ws = wb.active; ws.title = "Distributions"
    ws["A1"] = f"{tenant_name} — AP Distributions {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Fund", "Period", "Invoice #", "Account", "Amount"])
    r = 4
    summary: dict = defaultdict(Decimal)  # (fund, period) -> amount
    for d, inv, cc in rows:
        per = period_name(inv.gl_date)
        ws.cell(row=r, column=1, value=d.fund_value)
        ws.cell(row=r, column=2, value=per)
        ws.cell(row=r, column=3, value=inv.invoice_number)
        ws.cell(row=r, column=4, value=cc.concatenated_segments)
        ws.cell(row=r, column=5, value=float(d.amount))
        summary[(d.fund_value, per)] += Decimal(d.amount)
        r += 1
    for col, w in {"A": 12, "B": 12, "C": 16, "D": 30, "E": 14}.items():
        ws.column_dimensions[col].width = w
    # Summary sheet (Fund × Period)
    ss = wb.create_sheet("Summary by Fund-Period")
    _header_row(ss, 1, ["Fund", "Period", "Amount"])
    rr = 2
    for (fund, per), amt in sorted(summary.items()):
        ss.cell(row=rr, column=1, value=fund)
        ss.cell(row=rr, column=2, value=per)
        ss.cell(row=rr, column=3, value=float(amt))
        rr += 1
    for col, w in {"A": 12, "B": 12, "C": 16}.items():
        ss.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
