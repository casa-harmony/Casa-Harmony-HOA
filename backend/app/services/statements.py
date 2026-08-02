"""AR statement batch generation + email delivery with per-homeowner tracking."""
from __future__ import annotations

import io
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.statements import StatementDelivery, StatementRun
from app.models.subledger import ArHomeowner, ArInvoice, ArReceipt
from app.services import notifications
from app.services.reports import _BOLD, _header_row

_OPEN = ("DRAFT", "ACCOUNTED", "POSTED")


def _balance(db, tenant_id, homeowner_id) -> Decimal:
    rows = db.execute(select(ArInvoice.amount, ArInvoice.amount_paid).where(
        ArInvoice.tenant_id == tenant_id, ArInvoice.homeowner_id == homeowner_id,
        ArInvoice.status.in_(_OPEN))).all()
    return sum((Decimal(a) - Decimal(p or 0) for a, p in rows), Decimal("0"))


def build_statement_pdf(db, tenant_id, homeowner, as_of: date, tenant_name) -> bytes:
    open_invs = db.execute(select(ArInvoice).where(
        ArInvoice.tenant_id == tenant_id, ArInvoice.homeowner_id == homeowner.id,
        ArInvoice.status.in_(_OPEN)).order_by(ArInvoice.due_date)).scalars().all()
    receipts = db.execute(select(ArReceipt).where(
        ArReceipt.tenant_id == tenant_id, ArReceipt.homeowner_id == homeowner.id)
        .order_by(ArReceipt.receipt_date.desc()).limit(8)).scalars().all()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title=f"{tenant_name} Statement")
    st = getSampleStyleSheet()
    el = [Paragraph(f"{tenant_name} — Account Statement", st["Title"]),
          Paragraph(f"{homeowner.first_name} {homeowner.last_name} · Account {homeowner.account_number}", st["Normal"]),
          Paragraph(f"As of {as_of.isoformat()}", st["Normal"]), Spacer(1, 0.2 * inch)]
    bal = Decimal("0")
    rows = [["Invoice", "Type", "Due", "Balance"]]
    for i in open_invs:
        b = Decimal(i.amount) - Decimal(i.amount_paid or 0)
        bal += b
        rows.append([i.invoice_number, i.invoice_type, i.due_date.isoformat() if i.due_date else "—", f"${b:,.2f}"])
    if len(rows) == 1:
        rows.append(["—", "—", "—", "$0.00"])
    t = Table(rows, hAlign="LEFT", colWidths=[1.4 * inch, 1.6 * inch, 1.2 * inch, 1.2 * inch])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                           ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    el += [t, Spacer(1, 0.15 * inch),
           Paragraph(f"<b>Balance due: ${bal:,.2f}</b>", st["Normal"]), Spacer(1, 0.15 * inch)]
    if receipts:
        el.append(Paragraph("Recent payments", st["Heading3"]))
        for r in receipts:
            el.append(Paragraph(f"{r.receipt_date.isoformat()} — {r.receipt_number} — ${r.amount:,.2f} ({r.payment_method})", st["Normal"]))
    el += [Spacer(1, 0.2 * inch),
           Paragraph(f"Pay online at {settings.FRONTEND_BASE_URL}/portal", st["Normal"])]
    doc.build(el)
    return buf.getvalue()


def _next_run_number(db, tenant_id) -> str:
    n = db.execute(select(func.count(StatementRun.id)).where(StatementRun.tenant_id == tenant_id)).scalar_one()
    return f"STMT-{n + 1:05d}"


import hashlib
import hmac
import time

MAX_ATTACH_BYTES = 4_000_000  # attach inline below this; larger → signed link


def make_statement_link(tenant_id, homeowner_id, ttl_days: int = 30) -> str:
    """A signed, login-free link to download a homeowner's statement PDF."""
    exp = int(time.time()) + ttl_days * 86400
    msg = f"{tenant_id}:{homeowner_id}:{exp}".encode()
    sig = hmac.new(settings.SECRET_KEY.encode(), msg, hashlib.sha256).hexdigest()
    base = settings.FRONTEND_BASE_URL.rstrip("/")
    return f"{base}/api/v1/statements/public?t={tenant_id}&h={homeowner_id}&e={exp}&sig={sig}"


def verify_statement_link(tenant_id, homeowner_id, exp: int, sig: str) -> bool:
    if exp < int(time.time()):
        return False
    msg = f"{tenant_id}:{homeowner_id}:{exp}".encode()
    expected = hmac.new(settings.SECRET_KEY.encode(), msg, hashlib.sha256).hexdigest()
    return hmac.compare_digest(expected, sig)


def run_statements(db: Session, *, tenant_id, as_of: date, homeowner_ids=None,
                   send_email=True, attach_pdf=True, created_by=None) -> StatementRun:
    if homeowner_ids:
        homeowners = [h for h in (db.get(ArHomeowner, hid) for hid in homeowner_ids)
                      if h and h.tenant_id == tenant_id]
    else:
        homeowners = db.execute(select(ArHomeowner).where(
            ArHomeowner.tenant_id == tenant_id, ArHomeowner.status == "active")).scalars().all()

    run = StatementRun(tenant_id=tenant_id, run_number=_next_run_number(db, tenant_id),
                       as_of_date=as_of, status="COMPLETED",
                       created_by=created_by, updated_by=created_by)
    db.add(run)
    db.flush()

    tname = None
    from app.models.identity import Tenant
    t = db.get(Tenant, tenant_id)
    tname = t.name if t else "HOA"

    for ho in homeowners:
        bal = _balance(db, tenant_id, ho.id)
        pdf = build_statement_pdf(db, tenant_id, ho, as_of, tname)
        d = StatementDelivery(tenant_id=tenant_id, run_id=run.id, homeowner_id=ho.id,
                              email=ho.email, balance=bal, status="GENERATED",
                              created_by=created_by, updated_by=created_by)
        run.generated += 1
        if not send_email:
            d.status = "GENERATED"
        elif ho.statement_opt_out:
            d.status = "SKIPPED_OPTOUT"; run.skipped += 1
        elif not ho.email or ho.email.endswith("@anonymized.invalid"):
            d.status = "NO_EMAIL"; run.skipped += 1
        else:
            try:
                body = (f"Your balance as of {as_of.isoformat()} is ${bal:,.2f}. "
                        f"View details and pay online at {settings.FRONTEND_BASE_URL}/portal")
                attachments = None
                if attach_pdf and len(pdf) <= MAX_ATTACH_BYTES:
                    attachments = [(f"statement_{ho.account_number}.pdf", pdf, "application/pdf")]
                    d.attached = True
                elif attach_pdf:
                    # Too large to attach — provide a signed download link instead.
                    body += f"\n\nYour statement PDF: {make_statement_link(tenant_id, ho.id)}"
                notifications._send_email(ho.email, f"{tname} — your account statement", body,
                                          attachments=attachments)
                d.status = "SENT"; d.sent_at = datetime.now(timezone.utc); run.sent += 1
            except Exception as exc:  # pragma: no cover - provider errors recorded, not raised
                d.status = "FAILED"; d.error = str(exc)[:400]; run.failed += 1
        db.add(d)
    db.flush()
    return run


def build_delivery_report_workbook(db, tenant_id, run: StatementRun, tenant_name) -> bytes:
    rows = db.execute(
        select(StatementDelivery, ArHomeowner)
        .join(ArHomeowner, ArHomeowner.id == StatementDelivery.homeowner_id)
        .where(StatementDelivery.run_id == run.id)
        .order_by(ArHomeowner.account_number)).all()
    wb = Workbook(); ws = wb.active; ws.title = "Statement Delivery"
    ws["A1"] = f"{tenant_name} — Statement Run {run.run_number} ({run.as_of_date.isoformat()})"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Generated {run.generated} · Sent {run.sent} · Skipped {run.skipped} · Failed {run.failed}"
    _header_row(ws, 4, ["Account", "Homeowner", "Email", "Balance", "Status", "Attached", "Sent At"])
    r = 5
    for d, ho in rows:
        ws.cell(row=r, column=1, value=ho.account_number)
        ws.cell(row=r, column=2, value=f"{ho.first_name} {ho.last_name}")
        ws.cell(row=r, column=3, value=d.email or "")
        ws.cell(row=r, column=4, value=float(d.balance))
        ws.cell(row=r, column=5, value=d.status)
        ws.cell(row=r, column=6, value="Y" if d.attached else "")
        ws.cell(row=r, column=7, value=d.sent_at.isoformat() if d.sent_at else "")
        r += 1
    for col, w in {"A": 14, "B": 24, "C": 28, "D": 12, "E": 16, "F": 10, "G": 22}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
