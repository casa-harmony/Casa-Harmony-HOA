"""Purchasing reports: PO commitment register, variance by contract, cost-center
utilization, PO activity log, and matched-vs-nonmatched invoice summary."""
from __future__ import annotations

import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from app.models.audit import AuditLog
from app.models.kff import GlCodeCombination
from app.models.masters import ApSupplier
from app.models.payables import ApInvoice
from app.models.procurement import PoDistribution, PoHeader, PoLine
from app.services.reports import _BOLD, _header_row


def build_commitment_register(db, tenant_id, tenant_name) -> bytes:
    rows = db.execute(
        select(PoHeader, ApSupplier).join(ApSupplier, ApSupplier.id == PoHeader.vendor_id)
        .where(PoHeader.tenant_id == tenant_id).order_by(PoHeader.po_number)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Commitment Register"
    ws["A1"] = f"{tenant_name} — PO Commitment Register"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["PO #", "Type", "Vendor", "Start", "End", "Committed",
                        "Limit", "Billed", "Remaining", "Status"])
    r = 4
    tl = tb = Decimal("0")
    for po, v in rows:
        remaining = Decimal(po.amount_limit) - Decimal(po.billed_amount)
        ws.cell(row=r, column=1, value=po.po_number)
        ws.cell(row=r, column=2, value=po.document_type)
        ws.cell(row=r, column=3, value=v.name)
        ws.cell(row=r, column=4, value=po.start_date.isoformat() if po.start_date else "")
        ws.cell(row=r, column=5, value=po.end_date.isoformat() if po.end_date else "")
        ws.cell(row=r, column=6, value=float(po.amount))
        ws.cell(row=r, column=7, value=float(po.amount_limit))
        ws.cell(row=r, column=8, value=float(po.billed_amount))
        ws.cell(row=r, column=9, value=float(remaining))
        ws.cell(row=r, column=10, value=po.status)
        tl += Decimal(po.amount_limit); tb += Decimal(po.billed_amount)
        r += 1
    ws.cell(row=r, column=5, value="TOTAL").font = _BOLD
    ws.cell(row=r, column=7, value=float(tl)).font = _BOLD
    ws.cell(row=r, column=8, value=float(tb)).font = _BOLD
    ws.cell(row=r, column=9, value=float(tl - tb)).font = _BOLD
    for col, w in {"A": 14, "B": 10, "C": 26, "D": 12, "E": 12, "F": 13, "G": 13, "H": 13, "I": 13, "J": 16}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_variance_by_contract(db, tenant_id, tenant_name) -> bytes:
    rows = db.execute(
        select(PoHeader, ApSupplier).join(ApSupplier, ApSupplier.id == PoHeader.vendor_id)
        .where(PoHeader.tenant_id == tenant_id, PoHeader.document_type == "CONTRACT")
        .order_by(PoHeader.po_number)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Variance by Contract"
    ws["A1"] = f"{tenant_name} — Contract Variance (Ordered vs Billed)"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Contract", "Vendor", "Limit (Ordered)", "Billed", "Variance", "% Used"])
    r = 4
    for po, v in rows:
        limit = Decimal(po.amount_limit); billed = Decimal(po.billed_amount)
        pct = (billed / limit * 100) if limit else Decimal("0")
        ws.cell(row=r, column=1, value=po.po_number)
        ws.cell(row=r, column=2, value=v.name)
        ws.cell(row=r, column=3, value=float(limit))
        ws.cell(row=r, column=4, value=float(billed))
        ws.cell(row=r, column=5, value=float(limit - billed))
        ws.cell(row=r, column=6, value=round(float(pct), 1))
        r += 1
    for col, w in {"A": 14, "B": 26, "C": 16, "D": 14, "E": 14, "F": 10}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_cost_center_utilization(db, tenant_id, tenant_name) -> bytes:
    rows = db.execute(
        select(PoDistribution, GlCodeCombination)
        .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id)
        .where(PoDistribution.tenant_id == tenant_id)
    ).all()
    by_cc: dict = defaultdict(lambda: {"ordered": Decimal("0"), "billed": Decimal("0")})
    for d, cc in rows:
        key = cc.cost_center_value or "(none)"
        by_cc[key]["ordered"] += Decimal(d.amount)
        by_cc[key]["billed"] += Decimal(d.amount_billed)
    wb = Workbook(); ws = wb.active; ws.title = "Cost Center Utilization"
    ws["A1"] = f"{tenant_name} — Cost Center Utilization"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Cost Center", "Committed", "Billed", "Remaining", "% Used"])
    r = 4
    for key in sorted(by_cc):
        o = by_cc[key]["ordered"]; b = by_cc[key]["billed"]
        pct = (b / o * 100) if o else Decimal("0")
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=float(o))
        ws.cell(row=r, column=3, value=float(b))
        ws.cell(row=r, column=4, value=float(o - b))
        ws.cell(row=r, column=5, value=round(float(pct), 1))
        r += 1
    for col, w in {"A": 16, "B": 14, "C": 14, "D": 14, "E": 10}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_matched_summary(db, tenant_id, tenant_name) -> bytes:
    """Matched vs non-matched invoice summary (count + amount by match_status)."""
    rows = db.execute(
        select(ApInvoice).where(ApInvoice.tenant_id == tenant_id,
                                ApInvoice.status != "CANCELLED")
    ).scalars().all()
    agg: dict = defaultdict(lambda: {"count": 0, "amount": Decimal("0")})
    for inv in rows:
        agg[inv.match_status]["count"] += 1
        agg[inv.match_status]["amount"] += Decimal(inv.amount)
    wb = Workbook(); ws = wb.active; ws.title = "Matched vs Non-matched"
    ws["A1"] = f"{tenant_name} — Matched vs Non-matched Invoices"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Match Status", "Count", "Amount"])
    r = 4
    for k in sorted(agg):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=agg[k]["count"])
        ws.cell(row=r, column=3, value=float(agg[k]["amount"]))
        r += 1
    for col, w in {"A": 20, "B": 10, "C": 16}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_po_activity_log(db, tenant_id, tenant_name) -> bytes:
    """PO/invoice audit activity (creates, approvals, holds, reversals)."""
    rows = db.execute(
        select(AuditLog).where(
            AuditLog.tenant_id == tenant_id,
            AuditLog.entity_type.in_(("PoHeader", "ApInvoice", "ApMatchTolerance")))
        .order_by(AuditLog.created_at.desc()).limit(2000)
    ).scalars().all()
    wb = Workbook(); ws = wb.active; ws.title = "PO Activity Log"
    ws["A1"] = f"{tenant_name} — PO/AP Activity Log"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["When", "Entity", "Action", "Entity ID"])
    r = 4
    for a in rows:
        ws.cell(row=r, column=1, value=a.created_at.isoformat() if a.created_at else "")
        ws.cell(row=r, column=2, value=a.entity_type)
        ws.cell(row=r, column=3, value=a.action)
        ws.cell(row=r, column=4, value=str(a.entity_id) if a.entity_id else "")
        r += 1
    for col, w in {"A": 28, "B": 16, "C": 16, "D": 40}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_contract_utilization(db, tenant_id, tenant_name) -> bytes:
    """Contract utilization by Cost Center × Fund (committed vs billed)."""
    rows = db.execute(
        select(PoDistribution, GlCodeCombination)
        .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id)
        .join(PoLine, PoLine.id == PoDistribution.po_line_id)
        .join(PoHeader, PoHeader.id == PoLine.po_header_id)
        .where(PoDistribution.tenant_id == tenant_id,
               PoHeader.document_type == "CONTRACT")
    ).all()
    agg: dict = defaultdict(lambda: {"ordered": Decimal("0"), "billed": Decimal("0")})
    for d, cc in rows:
        key = (cc.cost_center_value or "(none)", d.fund_value)
        agg[key]["ordered"] += Decimal(d.amount)
        agg[key]["billed"] += Decimal(d.amount_billed)
    wb = Workbook(); ws = wb.active; ws.title = "Contract Utilization"
    ws["A1"] = f"{tenant_name} — Contract Utilization by Cost Center / Fund"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Cost Center", "Fund", "Committed", "Billed", "Remaining", "% Used"])
    r = 4
    for key in sorted(agg):
        o = agg[key]["ordered"]; b = agg[key]["billed"]
        pct = (b / o * 100) if o else Decimal("0")
        ws.cell(row=r, column=1, value=key[0])
        ws.cell(row=r, column=2, value=key[1])
        ws.cell(row=r, column=3, value=float(o))
        ws.cell(row=r, column=4, value=float(b))
        ws.cell(row=r, column=5, value=float(o - b))
        ws.cell(row=r, column=6, value=round(float(pct), 1))
        r += 1
    for col, w in {"A": 16, "B": 12, "C": 14, "D": 14, "E": 14, "F": 10}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_receiving_register_workbook(db, tenant_id, start, end, tenant_name) -> bytes:
    """Receiving register: accepted/pending receipt transactions by PO, cost center, Fund."""
    from app.models.receiving import RcvShipmentHeader, RcvShipmentLine, RcvTransaction

    rows = db.execute(
        select(RcvTransaction, RcvShipmentHeader, PoHeader, GlCodeCombination, ApSupplier)
        .join(RcvShipmentLine, RcvShipmentLine.id == RcvTransaction.shipment_line_id)
        .join(RcvShipmentHeader, RcvShipmentHeader.id == RcvShipmentLine.header_id)
        .join(PoHeader, PoHeader.id == RcvShipmentHeader.po_header_id)
        .join(ApSupplier, ApSupplier.id == PoHeader.vendor_id)
        .join(PoDistribution, PoDistribution.id == RcvTransaction.po_distribution_id)
        .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id)
        .where(RcvShipmentHeader.tenant_id == tenant_id,
               RcvShipmentHeader.received_date >= start,
               RcvShipmentHeader.received_date <= end)
        .order_by(RcvShipmentHeader.received_date)
    ).all()
    wb = Workbook(); ws = wb.active; ws.title = "Receiving Register"
    ws["A1"] = f"{tenant_name} — Receiving Register {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Date", "Receipt #", "PO #", "Vendor", "Cost Center", "Fund",
                        "Account", "Qty", "Amount", "Accepted"])
    r = 4
    by_fund: dict = defaultdict(Decimal)
    total = Decimal("0")
    for txn, hdr, po, cc, v in rows:
        ws.cell(row=r, column=1, value=hdr.received_date.isoformat())
        ws.cell(row=r, column=2, value=hdr.receipt_number)
        ws.cell(row=r, column=3, value=po.po_number)
        ws.cell(row=r, column=4, value=v.name)
        ws.cell(row=r, column=5, value=cc.cost_center_value or "")
        ws.cell(row=r, column=6, value=txn.fund_value)
        ws.cell(row=r, column=7, value=cc.concatenated_segments)
        ws.cell(row=r, column=8, value=float(txn.quantity))
        ws.cell(row=r, column=9, value=float(txn.amount))
        ws.cell(row=r, column=10, value="Y" if txn.accepted else "N")
        if txn.accepted:
            by_fund[txn.fund_value] += Decimal(txn.amount)
            total += Decimal(txn.amount)
        r += 1
    ws.cell(row=r, column=8, value="TOTAL (accepted)").font = _BOLD
    ws.cell(row=r, column=9, value=float(total)).font = _BOLD
    r += 2
    ws.cell(row=r, column=1, value="Accepted received by Fund").font = _BOLD
    r += 1
    for fund, amt in sorted(by_fund.items()):
        ws.cell(row=r, column=1, value=fund)
        ws.cell(row=r, column=2, value=float(amt))
        r += 1
    for col, w in {"A": 12, "B": 14, "C": 12, "D": 24, "E": 12, "F": 10, "G": 28, "H": 10, "I": 12, "J": 9}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def commitment_balances(db, tenant_id):
    """Open commitment by (cost center, fund) from PO distributions of live POs."""
    rows = db.execute(
        select(PoDistribution, GlCodeCombination)
        .join(PoLine, PoLine.id == PoDistribution.po_line_id)
        .join(PoHeader, PoHeader.id == PoLine.po_header_id)
        .join(GlCodeCombination, GlCodeCombination.id == PoDistribution.code_combination_id)
        .where(PoHeader.tenant_id == tenant_id,
               PoHeader.status.in_(("APPROVED", "PARTIALLY_BILLED", "FULLY_BILLED")))
    ).all()
    agg: dict = defaultdict(lambda: [Decimal("0"), Decimal("0")])  # (cc,fund)->[committed,billed]
    for d, cc in rows:
        key = (cc.cost_center_value or "—", d.fund_value)
        agg[key][0] += Decimal(d.amount)
        agg[key][1] += Decimal(d.amount_billed)
    out = []
    for (cost_center, fund), (committed, billed) in sorted(agg.items()):
        out.append({"cost_center": cost_center, "fund_value": fund, "committed": committed,
                    "billed": billed, "available": committed - billed})
    return out


def build_encumbrance_register(db, tenant_id, tenant_name) -> bytes:
    """Encumbrance / commitment register by Fund and Cost Center."""
    data = commitment_balances(db, tenant_id)
    wb = Workbook(); ws = wb.active; ws.title = "Encumbrance Register"
    ws["A1"] = f"{tenant_name} — Encumbrance / Commitment Register"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Fund", "Cost Center", "Committed", "Billed (Actual)", "Open Commitment"])
    r = 4
    tc = tb = Decimal("0")
    for row in sorted(data, key=lambda x: (x["fund_value"], x["cost_center"])):
        ws.cell(row=r, column=1, value=row["fund_value"])
        ws.cell(row=r, column=2, value=row["cost_center"])
        ws.cell(row=r, column=3, value=float(row["committed"]))
        ws.cell(row=r, column=4, value=float(row["billed"]))
        ws.cell(row=r, column=5, value=float(row["available"]))
        tc += row["committed"]; tb += row["billed"]
        r += 1
    ws.cell(row=r, column=2, value="TOTAL").font = _BOLD
    ws.cell(row=r, column=3, value=float(tc)).font = _BOLD
    ws.cell(row=r, column=4, value=float(tb)).font = _BOLD
    ws.cell(row=r, column=5, value=float(tc - tb)).font = _BOLD
    for col, w in {"A": 10, "B": 16, "C": 16, "D": 16, "E": 18}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
