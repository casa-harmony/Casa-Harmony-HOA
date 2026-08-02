"""Go-live readiness: automated system checks + manual checklist + health metrics
+ exportable go-live / compliance package (xlsx).
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.audit import AuditLog
from app.models.budgeting import BudgetControlSettings
from app.models.cash import CeStatementHeader
from app.models.fixed_assets import FaAsset
from app.models.gl import GlJeBatch
from app.models.golive import ComplianceItem
from app.models.identity import Tenant, User
from app.models.kff import KffStructure
from app.models.masters import ApSupplier
from app.models.payables import ApInvoice
from app.models.period import AccountingPeriod

# Manual go-live items seeded per tenant (admins mark DONE/NA).
DEFAULT_MANUAL = [
    ("rotate_cloud_keys", "Rotate any cloud/admin access keys shared during build", "Security"),
    ("rotate_repo_token", "Rotate the source-control (GitHub) access token", "Security"),
    ("rotate_app_secret", "Set a strong, unique SECRET_KEY / FERNET key in production", "Security"),
    ("configure_backups", "Configure automated database backups + restore test", "Operations"),
    ("review_user_access", "Review user accounts, roles, and MFA enrollment", "Security"),
    ("verify_tls", "Verify HTTPS/TLS and security headers on the public domain", "Security"),
    ("seed_retained_earnings", "Confirm a Retained Earnings (3000) account exists per Fund", "Accounting"),
]


def ensure_items(db: Session, tenant_id) -> None:
    existing = {c.code for c in db.execute(select(ComplianceItem).where(
        ComplianceItem.tenant_id == tenant_id)).scalars().all()}
    for code, title, cat in DEFAULT_MANUAL:
        if code not in existing:
            db.add(ComplianceItem(tenant_id=tenant_id, code=code, title=title,
                                  category=cat, status="PENDING"))
    db.flush()


def _count(db, stmt) -> int:
    return int(db.execute(stmt).scalar_one())


def run_health(db: Session, tenant_id) -> dict:
    posted = _count(db, select(func.count(GlJeBatch.id)).where(
        GlJeBatch.tenant_id == tenant_id, GlJeBatch.status == "POSTED"))
    unbalanced = _count(db, select(func.count(GlJeBatch.id)).where(
        GlJeBatch.tenant_id == tenant_id, GlJeBatch.status == "POSTED",
        GlJeBatch.control_total_dr != GlJeBatch.control_total_cr))
    open_periods = _count(db, select(func.count(AccountingPeriod.id)).where(
        AccountingPeriod.tenant_id == tenant_id, AccountingPeriod.status == "OPEN"))
    closed_periods = _count(db, select(func.count(AccountingPeriod.id)).where(
        AccountingPeriod.tenant_id == tenant_id, AccountingPeriod.status == "CLOSED"))
    bc = db.execute(select(BudgetControlSettings).where(
        BudgetControlSettings.tenant_id == tenant_id)).scalar_one_or_none()
    return {
        "gl_posted_batches": posted,
        "gl_unbalanced_batches": unbalanced,
        "gl_balanced": unbalanced == 0,
        "open_periods": open_periods,
        "closed_periods": closed_periods,
        "budget_control_mode": bc.mode if bc else "NONE",
        "ap_open_holds": _count(db, select(func.count(ApInvoice.id)).where(
            ApInvoice.tenant_id == tenant_id, ApInvoice.on_hold.is_(True))),
        "unreconciled_statements": _count(db, select(func.count(CeStatementHeader.id)).where(
            CeStatementHeader.tenant_id == tenant_id, CeStatementHeader.status == "OPEN")),
        "active_assets": _count(db, select(func.count(FaAsset.id)).where(
            FaAsset.tenant_id == tenant_id, FaAsset.status == "ACTIVE")),
        "coa_structures": _count(db, select(func.count(KffStructure.id)).where(
            KffStructure.tenant_id == tenant_id)),
        "vendors": _count(db, select(func.count(ApSupplier.id)).where(
            ApSupplier.tenant_id == tenant_id)),
        "audit_events": _count(db, select(func.count(AuditLog.id)).where(
            AuditLog.tenant_id == tenant_id)),
        "superadmin_exists": _count(db, select(func.count(User.id)).where(
            User.is_superadmin.is_(True))) > 0,
    }


def _auto_checks(h: dict) -> list[dict]:
    def chk(code, title, ok, detail, warn_only=False):
        status = "PASS" if ok else ("WARN" if warn_only else "FAIL")
        return {"code": code, "title": title, "category": "Automated",
                "status": status, "detail": detail, "manual": False}
    return [
        chk("gl_integrity", "GL integrity: all posted batches balanced",
            h["gl_balanced"], f"{h['gl_unbalanced_batches']} unbalanced of {h['gl_posted_batches']} posted"),
        chk("coa_configured", "Chart of Accounts configured",
            h["coa_structures"] >= 1, f"{h['coa_structures']} COA structure(s)"),
        chk("audit_logging", "Audit logging active",
            h["audit_events"] > 0, f"{h['audit_events']} audit events", warn_only=True),
        chk("superadmin", "Platform super-admin present", h["superadmin_exists"], "superadmin account"),
        chk("budgetary_control", "Budgetary control enabled",
            h["budget_control_mode"] != "NONE",
            f"mode = {h['budget_control_mode']}", warn_only=True),
        chk("period_controls", "Period close controls in use",
            (h["open_periods"] + h["closed_periods"]) > 0,
            f"{h['open_periods']} open / {h['closed_periods']} closed", warn_only=True),
        chk("ap_holds_clear", "No lingering AP invoice holds",
            h["ap_open_holds"] == 0, f"{h['ap_open_holds']} on hold", warn_only=True),
        chk("bank_reconciled", "Bank statements reconciled",
            h["unreconciled_statements"] == 0,
            f"{h['unreconciled_statements']} open statement(s)", warn_only=True),
    ]


def checklist(db: Session, tenant_id) -> dict:
    ensure_items(db, tenant_id)
    h = run_health(db, tenant_id)
    auto = _auto_checks(h)
    manual = []
    for c in db.execute(select(ComplianceItem).where(
            ComplianceItem.tenant_id == tenant_id).order_by(ComplianceItem.category)).scalars().all():
        status = {"DONE": "PASS", "NA": "INFO", "PENDING": "WARN"}.get(c.status, "WARN")
        manual.append({"code": c.code, "title": c.title, "category": c.category,
                       "status": status, "detail": c.notes or "", "manual": True})
    items = auto + manual
    summary = {s: sum(1 for i in items if i["status"] == s) for s in ("PASS", "WARN", "FAIL", "INFO")}
    ready = summary["FAIL"] == 0 and all(
        m["status"] != "WARN" for m in manual if m["code"] in
        ("rotate_cloud_keys", "rotate_repo_token", "rotate_app_secret"))
    return {"health": h, "items": items, "summary": summary, "go_live_ready": ready}


def set_item(db: Session, tenant_id, code: str, status: str, notes: str | None = None) -> ComplianceItem:
    ensure_items(db, tenant_id)
    c = db.execute(select(ComplianceItem).where(
        ComplianceItem.tenant_id == tenant_id, ComplianceItem.code == code)).scalar_one_or_none()
    if c is None:
        raise ValueError("Checklist item not found")
    c.status = status
    if notes is not None:
        c.notes = notes
    db.flush()
    return c


def build_golive_package(db: Session, tenant_id, tenant_name: str) -> bytes:
    """Multi-sheet go-live / compliance package: checklist, health, readiness summary."""
    data = checklist(db, tenant_id)
    h = data["health"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Checklist"
    ws["A1"] = f"{tenant_name} — Go-Live Checklist"
    ws["A1"].font = Font(size=14, bold=True)
    hdr = ["Category", "Item", "Status", "Detail", "Type"]
    for i, t in enumerate(hdr, start=1):
        ws.cell(row=3, column=i, value=t).font = Font(bold=True)
    r = 4
    for it in data["items"]:
        ws.cell(row=r, column=1, value=it["category"])
        ws.cell(row=r, column=2, value=it["title"])
        ws.cell(row=r, column=3, value=it["status"])
        ws.cell(row=r, column=4, value=it["detail"])
        ws.cell(row=r, column=5, value="Manual" if it["manual"] else "Automated")
        r += 1
    for col, w in {"A": 14, "B": 48, "C": 8, "D": 40, "E": 12}.items():
        ws.column_dimensions[col].width = w

    hs = wb.create_sheet("Health")
    hs["A1"] = f"{tenant_name} — System Health"
    hs["A1"].font = Font(size=14, bold=True)
    rr = 3
    for k, v in h.items():
        hs.cell(row=rr, column=1, value=k)
        hs.cell(row=rr, column=2, value=str(v))
        rr += 1
    hs.column_dimensions["A"].width = 28
    hs.column_dimensions["B"].width = 16

    cs = wb.create_sheet("Readiness")
    cs["A1"] = "SOC 2 / PCI DSS Readiness Summary"
    cs["A1"].font = Font(size=14, bold=True)
    notes = [
        ("Tenant isolation", "PostgreSQL Row-Level Security on all tenant tables"),
        ("Encryption at rest", "Field-level encryption (Fernet) for bank/account numbers; PCI tokenization"),
        ("Access control", "RBAC with least-privilege roles; MFA (TOTP staff, email/SMS residents)"),
        ("Audit trail", f"Immutable audit log — {h['audit_events']} events recorded"),
        ("Segregation of duties", "Approval workflows for PO/AP/GL; period close lock-down"),
        ("GL integrity", "Balanced double-entry; " + ("BALANCED" if h["gl_balanced"] else "REVIEW UNBALANCED BATCHES")),
        ("Go-live ready", "YES" if data["go_live_ready"] else "NO — complete required security items"),
    ]
    rr = 3
    for k, v in notes:
        cs.cell(row=rr, column=1, value=k).font = Font(bold=True)
        cs.cell(row=rr, column=2, value=v)
        rr += 1
    cs.column_dimensions["A"].width = 24
    cs.column_dimensions["B"].width = 70

    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
