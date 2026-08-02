"""COA structure report export (xlsx via openpyxl)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.kff import (
    GlCodeCombination,
    KffSegment,
    KffStructure,
    KffValueSet,
    KffValueSetValue,
)

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="111827")


def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _autofit(ws, ncols: int, max_width: int = 60) -> None:
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        longest = max(
            (len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(12, longest + 2), max_width)


def build_coa_workbook(db: Session, structure: KffStructure, tenant_name: str) -> bytes:
    wb = Workbook()

    # --- Sheet 1: Structure & Segments ---
    ws = wb.active
    ws.title = "COA Structure"
    ws["A1"] = f"Casa Harmony AI — Chart of Accounts ({tenant_name})"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Structure: {structure.structure_code} — {structure.title}"
    ws["A3"] = f"Separator: '{structure.segment_separator}'   Currency: USD (single functional)"

    headers = [
        "Seg #", "Segment Name", "Prompt", "Column", "Qualifier",
        "Value Set", "Required", "Enabled", "Displayed", "Default",
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    _style_header(ws, header_row, len(headers))

    segments = db.execute(
        select(KffSegment)
        .where(KffSegment.structure_id == structure.id)
        .order_by(KffSegment.segment_number)
    ).scalars().all()
    for seg in segments:
        ws.append([
            seg.segment_number, seg.name, seg.prompt, seg.column_name, seg.qualifier,
            seg.value_set.code if seg.value_set else "(free-form)",
            "Yes" if seg.required else "No",
            "Yes" if seg.enabled else "No",
            "Yes" if seg.displayed else "No",
            seg.default_value or "",
        ])
    _autofit(ws, len(headers))

    # --- Sheet 2: Value Sets & Values ---
    ws2 = wb.create_sheet("Value Sets")
    vheaders = [
        "Value Set", "Validation", "Format", "Value", "Description",
        "Account Type", "Summary", "Posting", "Enabled",
    ]
    ws2.append(vheaders)
    _style_header(ws2, 1, len(vheaders))
    value_sets = db.execute(select(KffValueSet)).scalars().all()
    for vs in value_sets:
        values = db.execute(
            select(KffValueSetValue)
            .where(KffValueSetValue.value_set_id == vs.id)
            .order_by(KffValueSetValue.value)
        ).scalars().all()
        if not values:
            ws2.append([vs.code, vs.validation_type, vs.format_type, "", "(no values)", "", "", "", ""])
        for v in values:
            ws2.append([
                vs.code, vs.validation_type, vs.format_type, v.value, v.description or "",
                v.account_type or "", "Yes" if v.summary_flag else "No",
                "Yes" if v.allow_posting else "No", "Yes" if v.enabled else "No",
            ])
    _autofit(ws2, len(vheaders))

    # --- Sheet 3: Code Combinations ---
    ws3 = wb.create_sheet("Code Combinations")
    cheaders = [
        "Concatenated", "Balancing", "Fund", "Cost Center", "Natural Acct",
        "Acct Type", "Posting", "Enabled",
    ]
    ws3.append(cheaders)
    _style_header(ws3, 1, len(cheaders))
    combos = db.execute(
        select(GlCodeCombination)
        .where(GlCodeCombination.structure_id == structure.id)
        .order_by(GlCodeCombination.concatenated_segments)
    ).scalars().all()
    for cc in combos:
        ws3.append([
            cc.concatenated_segments, cc.balancing_segment_value or "", cc.fund_value or "",
            cc.cost_center_value or "", cc.natural_account_value or "", cc.account_type or "",
            "Yes" if cc.allow_posting else "No", "Yes" if cc.enabled else "No",
        ])
    _autofit(ws3, len(cheaders))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
