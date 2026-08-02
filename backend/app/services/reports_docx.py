"""Word (docx) financial statement exports (fund-based)."""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import Session

from app.services.reports import _TYPE_LABEL, _fund_statement_data


def _money(v: Decimal) -> str:
    return f"${v:,.2f}" if v >= 0 else f"(${abs(v):,.2f})"


def build_financial_statements_docx(
    db: Session, tenant_id: uuid.UUID, period_name: str, tenant_name: str
) -> bytes:
    data = _fund_statement_data(db, tenant_id, period_name)
    funds = sorted(data.keys())
    doc = Document()

    title = doc.add_heading(tenant_name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"Fund-Based Financial Statements — {period_name}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Functional currency: USD").alignment = WD_ALIGN_PARAGRAPH.CENTER

    def section(heading: str, types: list[str]) -> Decimal:
        doc.add_heading(heading, level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Fund", "Type", "Account", "Balance"
        total = Decimal("0")
        for fund in funds:
            for at in types:
                for acct, bal in data.get(fund, {}).get(at, []):
                    row = table.add_row().cells
                    row[0].text, row[1].text = fund, _TYPE_LABEL[at]
                    row[2].text, row[3].text = acct, _money(bal)
                    total += bal
        p = doc.add_paragraph()
        run = p.add_run(f"Total {heading}: {_money(total)}")
        run.bold = True
        run.font.size = Pt(11)
        return total

    doc.add_heading("Balance Sheet", level=1)
    assets = section("Assets", ["A"])
    liabilities = section("Liabilities", ["L"])
    equity = section("Fund Balance / Equity", ["O"])
    doc.add_paragraph(
        f"Assets {_money(assets)} = Liabilities {_money(liabilities)} + "
        f"Equity {_money(equity)}"
    ).runs[0].bold = True

    doc.add_page_break()
    doc.add_heading("Statement of Revenues & Expenses", level=1)
    revenue = section("Revenue", ["R"])
    expenses = section("Expenses", ["E"])
    net = doc.add_paragraph()
    nr = net.add_run(f"Net Surplus / (Deficit): {_money(revenue - expenses)}")
    nr.bold = True
    nr.font.size = Pt(12)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
