"""Dunning automation: evaluate the rule ladder against each delinquent homeowner,
send reminder emails (with statement PDF) or escalate the delinquency case, and log
every action. Deduped to once per (homeowner, rule) per cycle.
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.dunning import DunningLog, DunningRule
from app.models.subledger import ArHomeowner, ArInvoice
from app.services import collections, notifications, statements
from app.services.reports import _BOLD, _header_row

_OPEN = ("DRAFT", "ACCOUNTED", "POSTED")
CYCLE_DAYS = 25  # don't re-fire the same rule for a homeowner within this window


def _oldest_due_dpd(db, tenant_id, homeowner_id, as_of: date) -> tuple[Decimal, int]:
    rows = db.execute(select(ArInvoice).where(
        ArInvoice.tenant_id == tenant_id, ArInvoice.homeowner_id == homeowner_id,
        ArInvoice.status.in_(_OPEN))).scalars().all()
    bal = Decimal("0")
    oldest_due = None
    for i in rows:
        b = Decimal(i.amount) - Decimal(i.amount_paid or 0)
        if b <= 0:
            continue
        bal += b
        if i.due_date and i.due_date < as_of and (oldest_due is None or i.due_date < oldest_due):
            oldest_due = i.due_date
    dpd = (as_of - oldest_due).days if oldest_due else 0
    return bal, dpd


def _recent_log(db, tenant_id, homeowner_id, rule_id, as_of: date) -> bool:
    cutoff = as_of - timedelta(days=CYCLE_DAYS)
    return db.execute(select(DunningLog).where(
        DunningLog.tenant_id == tenant_id, DunningLog.homeowner_id == homeowner_id,
        DunningLog.rule_id == rule_id, DunningLog.created_at >= datetime(cutoff.year, cutoff.month, cutoff.day, tzinfo=timezone.utc))
    ).first() is not None


def run_dunning(db: Session, *, tenant_id, as_of: date, homeowner_ids=None, created_by=None) -> dict:
    rules = db.execute(select(DunningRule).where(
        DunningRule.tenant_id == tenant_id, DunningRule.active.is_(True))
        .order_by(DunningRule.days_past_due)).scalars().all()
    if not rules:
        return {"reminders_sent": 0, "escalations": 0, "skipped": 0}
    if homeowner_ids:
        homeowners = [h for h in (db.get(ArHomeowner, hid) for hid in homeowner_ids)
                      if h and h.tenant_id == tenant_id and h.status == "active"]
    else:
        homeowners = db.execute(select(ArHomeowner).where(
            ArHomeowner.tenant_id == tenant_id, ArHomeowner.status == "active")).scalars().all()
    from app.models.identity import Tenant
    t = db.get(Tenant, tenant_id)
    tname = t.name if t else "HOA"

    reminders = escalations = skipped = 0
    for ho in homeowners:
        bal, dpd = _oldest_due_dpd(db, tenant_id, ho.id, as_of)
        if bal <= 0 or dpd <= 0:
            continue
        # Fire the highest-threshold rule the account currently qualifies for.
        eligible = [r for r in rules if dpd >= r.days_past_due]
        if not eligible:
            continue
        rule = eligible[-1]
        if _recent_log(db, tenant_id, ho.id, rule.id, as_of):
            continue
        log = DunningLog(tenant_id=tenant_id, homeowner_id=ho.id, rule_id=rule.id,
                         days_past_due=dpd, action=rule.action, balance=int(bal),
                         created_by=created_by, updated_by=created_by)
        if rule.action == "ESCALATE":
            case = collections.open_case(db, tenant_id, ho.id, as_of, created_by=created_by)
            if rule.escalate_to_stage and case.stage != rule.escalate_to_stage:
                try:
                    collections.escalate(db, case, rule.escalate_to_stage)
                except collections.CollectionsError:
                    pass  # already at/past this stage — leave as-is
            log.status = "ESCALATED"; log.detail = f"Case → {case.stage}"; escalations += 1
        else:  # REMINDER
            if ho.statement_opt_out:
                log.status = "SKIPPED_OPTOUT"; skipped += 1
            elif not ho.email or ho.email.endswith("@anonymized.invalid"):
                log.status = "NO_EMAIL"; skipped += 1
            else:
                try:
                    body = (rule.message or
                            f"Your account is {dpd} days past due with a balance of ${bal:,.2f}. "
                            f"Please pay at {settings.FRONTEND_BASE_URL}/portal.")
                    attachments = None
                    if rule.attach_statement:
                        pdf = statements.build_statement_pdf(db, tenant_id, ho, as_of, tname)
                        attachments = [(f"statement_{ho.account_number}.pdf", pdf, "application/pdf")]
                    notifications._send_email(ho.email, f"{tname} — payment reminder", body,
                                              attachments=attachments)
                    log.status = "SENT"; log.sent_at = datetime.now(timezone.utc); reminders += 1
                except Exception as exc:  # pragma: no cover
                    log.status = "FAILED"; log.detail = str(exc)[:300]
        db.add(log)
    db.flush()
    return {"reminders_sent": reminders, "escalations": escalations, "skipped": skipped}


def recent_logs(db, tenant_id, limit=100):
    return db.execute(select(DunningLog).where(DunningLog.tenant_id == tenant_id)
                      .order_by(DunningLog.created_at.desc()).limit(limit)).scalars().all()


def effectiveness(db, tenant_id, start: date, end: date) -> dict:
    """Reminders sent in range + how many of those accounts are now current (cured)."""
    start_dt = datetime(start.year, start.month, start.day, tzinfo=timezone.utc)
    end_dt = datetime(end.year, end.month, end.day, 23, 59, 59, tzinfo=timezone.utc)
    logs = db.execute(select(DunningLog).where(
        DunningLog.tenant_id == tenant_id, DunningLog.action == "REMINDER",
        DunningLog.status == "SENT",
        DunningLog.created_at >= start_dt, DunningLog.created_at <= end_dt)).scalars().all()
    sent = len(logs)
    cured = 0
    for lg in logs:
        bal, _ = _oldest_due_dpd(db, tenant_id, lg.homeowner_id, date.today())
        if bal <= 0:
            cured += 1
    rate = round(cured / sent * 100, 1) if sent else 0.0
    return {"reminders_sent": sent, "cured": cured, "cure_rate_pct": rate,
            "escalations": db.execute(select(func.count(DunningLog.id)).where(
                DunningLog.tenant_id == tenant_id, DunningLog.action == "ESCALATE",
                DunningLog.created_at >= start_dt, DunningLog.created_at <= end_dt)).scalar_one()}


def build_effectiveness_workbook(db, tenant_id, start: date, end: date, tenant_name) -> bytes:
    eff = effectiveness(db, tenant_id, start, end)
    wb = Workbook(); ws = wb.active; ws.title = "Dunning Effectiveness"
    ws["A1"] = f"{tenant_name} — Dunning Effectiveness {start.isoformat()} to {end.isoformat()}"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Metric", "Value"])
    rows = [("Reminders sent", eff["reminders_sent"]), ("Accounts cured", eff["cured"]),
            ("Cure rate %", eff["cure_rate_pct"]), ("Escalations", eff["escalations"])]
    r = 4
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = _BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
