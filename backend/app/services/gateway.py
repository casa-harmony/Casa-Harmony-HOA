"""Payment gateway service — provider-abstracted online collections.

Flow: create_checkout (PENDING txn) -> provider hosted checkout -> async webhook
confirms -> AR receipt + GL posting (idempotent). Refunds post a reversing draft
GL batch and restore the invoice balance. The MOCK provider lets the full flow run
end-to-end without external calls; a Stripe-like adapter slots into the same seams.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.payment_gateway import GatewayConfig, GatewayTransaction
from app.models.gl import GlJeLine
from app.models.subledger import ArInvoice, ArReceipt
from app.services import subledger_accounting as sla
from app.services.reports import _BOLD, _header_row

CENT = Decimal("0.01")
AR_RECEIVABLE = "1100"


class GatewayError(ValueError):
    pass


def get_config(db: Session, tenant_id) -> GatewayConfig:
    c = db.execute(select(GatewayConfig).where(GatewayConfig.tenant_id == tenant_id)).scalar_one_or_none()
    if c is None:
        from app.core.model_defaults import make_default
        c = make_default(GatewayConfig, tenant_id=tenant_id)
    return c


def _cash_natural(fund: str) -> str:
    return "1010" if fund == "RESV" else "1000"


def create_checkout(db: Session, *, tenant_id, invoice_id, amount, homeowner_id=None, created_by=None) -> dict:
    from app.services import gateway_providers as gp
    cfg = get_config(db, tenant_id)
    if not cfg.active:
        raise GatewayError("Payment gateway is not enabled")
    inv = db.get(ArInvoice, invoice_id)
    if inv is None or inv.tenant_id != tenant_id:
        raise GatewayError("Invoice not found")
    amt = Decimal(str(amount)).quantize(CENT)
    if amt <= 0:
        raise GatewayError("Amount must be positive")

    base = settings.FRONTEND_BASE_URL.rstrip("/")
    if cfg.provider == "STRIPE":
        try:
            res = gp.stripe_create_checkout(
                secret_key=cfg.secret_key, amount=amt,
                description=f"{inv.invoice_number} payment",
                success_url=f"{base}/portal?paid={inv.id}", cancel_url=f"{base}/portal",
                metadata={"invoice_id": str(inv.id)})
        except gp.ProviderError as exc:
            raise GatewayError(str(exc))
        txn_ref, checkout_url, client_secret = res["txn_ref"], res["checkout_url"], res.get("client_secret")
    else:  # MOCK
        txn_ref = f"mock_{uuid.uuid4().hex[:24]}"
        res = gp.mock_checkout(txn_ref, amt)
        checkout_url, client_secret = res["checkout_url"], res["client_secret"]

    txn = GatewayTransaction(
        tenant_id=tenant_id, txn_ref=txn_ref, provider=cfg.provider,
        homeowner_id=homeowner_id or inv.homeowner_id, invoice_id=inv.id, amount=amt,
        status="PENDING", created_by=created_by, updated_by=created_by)
    db.add(txn)
    db.flush()
    return {"txn_ref": txn_ref, "checkout_url": checkout_url, "client_secret": client_secret,
            "status": "PENDING", "amount": str(amt), "provider": cfg.provider}


def confirm_payment(db: Session, tenant_id, txn_ref: str) -> GatewayTransaction:
    """Idempotently confirm a successful payment: create the AR receipt + post GL."""
    txn = db.execute(select(GatewayTransaction).where(
        GatewayTransaction.tenant_id == tenant_id,
        GatewayTransaction.txn_ref == txn_ref)).scalar_one_or_none()
    if txn is None:
        raise GatewayError("Transaction not found")
    if txn.status == "SUCCEEDED":
        return txn  # idempotent — already applied
    inv = db.get(ArInvoice, txn.invoice_id) if txn.invoice_id else None
    seq = db.execute(select(func.count(ArReceipt.id)).where(ArReceipt.tenant_id == tenant_id)).scalar_one()
    receipt = ArReceipt(
        tenant_id=tenant_id, homeowner_id=txn.homeowner_id, applied_invoice_id=txn.invoice_id,
        receipt_number=f"GW-{seq + 1:06d}", amount=txn.amount, receipt_date=date.today(),
        payment_method="CARD", status="APPLIED", created_by=None, updated_by=None)
    db.add(receipt)
    db.flush()
    fund = (inv.fund if inv else None) or "OPER"
    if inv is not None:
        inv.amount_paid = (Decimal(inv.amount_paid or 0) + Decimal(txn.amount))
        if inv.amount_paid >= Decimal(inv.amount):
            inv.status = "PAID"
    from app.services.subledger_accounting import create_accounting_for_ar_receipt
    create_accounting_for_ar_receipt(db, receipt, fund=fund)
    txn.status = "SUCCEEDED"
    txn.receipt_id = receipt.id
    txn.confirmed_at = datetime.now(timezone.utc)
    db.flush()
    return txn


def _dispatch_event(db, tenant_id, event_type: str, txn_ref: str | None) -> dict:
    if event_type in ("payment_intent.succeeded", "checkout.session.completed", "payment_succeeded"):
        if not txn_ref:
            raise GatewayError("Missing transaction reference")
        txn = confirm_payment(db, tenant_id, txn_ref)
        return {"handled": True, "status": txn.status}
    if event_type in ("charge.refunded", "payment_refunded", "charge.dispute.created"):
        txn = db.execute(select(GatewayTransaction).where(
            GatewayTransaction.tenant_id == tenant_id,
            GatewayTransaction.txn_ref == txn_ref)).scalar_one_or_none()
        if txn and txn.status == "SUCCEEDED":
            t = refund(db, tenant_id=tenant_id, txn_id=txn.id)
            if event_type == "charge.dispute.created":
                t.status = "CHARGEBACK"
                db.flush()
        return {"handled": True, "status": "REFUNDED"}
    return {"handled": False, "status": "IGNORED"}


def handle_webhook_raw(db: Session, tenant_id, *, raw_body: bytes, sig_header: str | None,
                       json_signature: str | None = None) -> dict:
    """Verify + dispatch a provider webhook. STRIPE uses Stripe-Signature HMAC over the
    raw body; MOCK uses a shared-secret in the JSON/header and a {event_type,txn_ref} body."""
    import json as _json
    from app.services import gateway_providers as gp

    cfg = get_config(db, tenant_id)
    if cfg.provider == "STRIPE":
        if not gp.stripe_verify_signature(webhook_secret=cfg.webhook_secret or "",
                                          payload=raw_body, sig_header=sig_header or ""):
            raise GatewayError("Invalid webhook signature")
        event_type, txn_ref = gp.parse_stripe_event(raw_body)
        return _dispatch_event(db, tenant_id, event_type, txn_ref)
    # MOCK: shared-secret check against the JSON body fields.
    #
    # The secret must be explicitly configured. Never fall back to a guessable
    # constant like "mock" — with an unset secret that turned an unauthenticated
    # endpoint into a free "mark this invoice paid" button. A MOCK gateway with
    # no webhook secret cannot confirm payments; use a real provider (STRIPE)
    # with a configured secret to accept live webhooks.
    import hmac

    if settings.is_production:
        raise GatewayError("Mock gateway webhooks are disabled in production")
    if not cfg.webhook_secret:
        raise GatewayError("Webhook secret is not configured for this gateway")

    body = _json.loads(raw_body.decode("utf-8")) if raw_body else {}
    provided = json_signature or body.get("signature") or ""
    if not hmac.compare_digest(str(provided), str(cfg.webhook_secret)):
        raise GatewayError("Invalid webhook signature")
    return _dispatch_event(db, tenant_id, body.get("event_type", ""), body.get("txn_ref"))


def handle_webhook(db: Session, tenant_id, *, event_type: str, txn_ref: str, signature: str | None) -> dict:
    import hmac

    cfg = get_config(db, tenant_id)
    # Never default the secret to a guessable constant — see handle_webhook_raw.
    if not cfg.webhook_secret:
        raise GatewayError("Webhook secret is not configured for this gateway")
    if not signature or not hmac.compare_digest(str(signature), str(cfg.webhook_secret)):
        raise GatewayError("Invalid webhook signature")
    if event_type in ("payment_intent.succeeded", "checkout.session.completed", "payment_succeeded"):
        txn = confirm_payment(db, tenant_id, txn_ref)
        return {"handled": True, "status": txn.status}
    if event_type in ("charge.refunded", "payment_refunded"):
        txn = db.execute(select(GatewayTransaction).where(
            GatewayTransaction.tenant_id == tenant_id,
            GatewayTransaction.txn_ref == txn_ref)).scalar_one_or_none()
        if txn and txn.status == "SUCCEEDED":
            refund(db, tenant_id=tenant_id, txn_id=txn.id)
        return {"handled": True, "status": "REFUNDED"}
    return {"handled": False, "status": "IGNORED"}


def refund(db: Session, *, tenant_id, txn_id, created_by=None) -> GatewayTransaction:
    txn = db.get(GatewayTransaction, txn_id)
    if txn is None or txn.tenant_id != tenant_id:
        raise GatewayError("Transaction not found")
    if txn.status != "SUCCEEDED":
        raise GatewayError(f"Cannot refund a {txn.status} transaction")
    inv = db.get(ArInvoice, txn.invoice_id) if txn.invoice_id else None
    fund = (inv.fund if inv else None) or "OPER"
    structure = sla.get_primary_structure(db, tenant_id)
    recv = sla._account(db, tenant_id, structure.id, AR_RECEIVABLE, fund)
    cash = sla._account(db, tenant_id, structure.id, _cash_natural(fund), fund)
    batch = sla._new_batch(db, tenant_id, "AR", date.today(), f"Refund {txn.txn_ref[:18]}", created_by)
    header = sla._add_header(db, batch, structure.id, "Refund", "Gateway",
                             date.today(), "GW_REFUND", txn.id, f"Refund {txn.txn_ref[:18]}", created_by)
    amt = Decimal(txn.amount)
    # Reverse the original receipt: Dr AR / Cr Cash.
    db.add(GlJeLine(tenant_id=tenant_id, header_id=header.id, line_num=1,
                    code_combination_id=recv.id, entered_dr=amt, entered_cr=0, fund_value=fund,
                    description="Refund — restore receivable", created_by=created_by, updated_by=created_by))
    db.add(GlJeLine(tenant_id=tenant_id, header_id=header.id, line_num=2,
                    code_combination_id=cash.id, entered_dr=0, entered_cr=amt, fund_value=fund,
                    description="Refund — cash out", created_by=created_by, updated_by=created_by))
    db.flush()
    db.refresh(batch)
    sla._set_control_totals(batch)
    if inv is not None:
        inv.amount_paid = max(Decimal("0"), Decimal(inv.amount_paid or 0) - amt)
        if inv.status == "PAID" and inv.amount_paid < Decimal(inv.amount):
            inv.status = "POSTED"
    txn.status = "REFUNDED"
    txn.gl_je_header_id = header.id
    db.add(GatewayTransaction(
        tenant_id=tenant_id, txn_ref=f"refund_{txn.txn_ref}", provider=txn.provider,
        homeowner_id=txn.homeowner_id, invoice_id=txn.invoice_id, amount=amt, status="REFUNDED",
        refund_of_id=txn.id, gl_je_header_id=header.id, created_by=created_by, updated_by=created_by))
    db.flush()
    return txn


def build_reconciliation_workbook(db, tenant_id, tenant_name) -> bytes:
    rows = db.execute(select(GatewayTransaction).where(
        GatewayTransaction.tenant_id == tenant_id)
        .order_by(GatewayTransaction.created_at.desc())).scalars().all()
    wb = Workbook(); ws = wb.active; ws.title = "Gateway Reconciliation"
    ws["A1"] = f"{tenant_name} — Payment Gateway Reconciliation"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Txn Ref", "Provider", "Amount", "Status", "Receipt?", "When"])
    r = 4
    succeeded = Decimal("0")
    for t in rows:
        ws.cell(row=r, column=1, value=t.txn_ref)
        ws.cell(row=r, column=2, value=t.provider)
        ws.cell(row=r, column=3, value=float(t.amount))
        ws.cell(row=r, column=4, value=t.status)
        ws.cell(row=r, column=5, value="Y" if t.receipt_id else "")
        ws.cell(row=r, column=6, value=str(t.confirmed_at or t.created_at))
        if t.status == "SUCCEEDED":
            succeeded += Decimal(t.amount)
        r += 1
    ws.cell(row=r, column=2, value="TOTAL succeeded").font = _BOLD
    ws.cell(row=r, column=3, value=float(succeeded)).font = _BOLD
    for col, w in {"A": 34, "B": 10, "C": 12, "D": 12, "E": 9, "F": 24}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
