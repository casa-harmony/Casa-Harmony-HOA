"""Board & executive reporting: delinquency packet (PDF), cash-flow forecast (xlsx),
and an executive dashboard (JSON). Fund-aware throughout.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.budgeting import BudgetControlSettings, BudgetLine, BudgetVersion
from app.models.gl import GlBalance
from app.models.kff import GlCodeCombination
from app.models.subledger import ArInvoice
from app.services.collections import BUCKETS, aging, aging_by_fund, collection_effectiveness
from app.services.reports import _BOLD, _header_row

CASH_NATURALS = ("1000", "1010")
_OPEN = ("DRAFT", "ACCOUNTED", "POSTED")
_MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def _cash_by_fund(db: Session, tenant_id) -> dict:
    rows = db.execute(
        select(GlBalance, GlCodeCombination)
        .join(GlCodeCombination, GlCodeCombination.id == GlBalance.code_combination_id)
        .where(GlBalance.tenant_id == tenant_id,
               GlCodeCombination.natural_account_value.in_(CASH_NATURALS))
    ).all()
    out: dict = defaultdict(Decimal)
    for bal, cc in rows:
        out[cc.fund_value or "OPER"] += Decimal(bal.period_net_dr) - Decimal(bal.period_net_cr)
    return out


def _ar_open_by_fund(db: Session, tenant_id) -> dict:
    out: dict = defaultdict(Decimal)
    for inv in db.execute(select(ArInvoice).where(
            ArInvoice.tenant_id == tenant_id, ArInvoice.status.in_(_OPEN))).scalars():
        bal = Decimal(inv.amount) - Decimal(inv.amount_paid or 0)
        if bal > 0:
            out[inv.fund or "OPER"] += bal
    return out


def _controlling_budget(db, tenant_id):
    s = db.execute(select(BudgetControlSettings).where(
        BudgetControlSettings.tenant_id == tenant_id)).scalar_one_or_none()
    if s and s.controlling_version_id:
        return s.controlling_version_id
    v = db.execute(select(BudgetVersion).where(BudgetVersion.tenant_id == tenant_id)
                   .order_by(BudgetVersion.fiscal_year.desc())).scalars().first()
    return v.id if v else None


def cash_flow_forecast(db: Session, tenant_id, start: date, months: int = 6) -> list[dict]:
    """Rolling per-fund forecast: opening cash + AR due (inflow) − budgeted expense (outflow)."""
    cash = _cash_by_fund(db, tenant_id)
    funds = sorted(set(cash) | {"OPER", "RESV"})

    # Expected inflows: open AR balances by (fund, due year-month).
    inflow: dict = defaultdict(Decimal)
    for inv in db.execute(select(ArInvoice).where(
            ArInvoice.tenant_id == tenant_id, ArInvoice.status.in_(_OPEN))).scalars():
        bal = Decimal(inv.amount) - Decimal(inv.amount_paid or 0)
        if bal > 0 and inv.due_date:
            inflow[(inv.fund or "OPER", inv.due_date.year, inv.due_date.month)] += bal

    # Budgeted expense outflows by (fund, period_num) from the controlling/most-recent version.
    outflow: dict = defaultdict(Decimal)
    vid = _controlling_budget(db, tenant_id)
    if vid:
        for line, cc in db.execute(
            select(BudgetLine, GlCodeCombination)
            .join(GlCodeCombination, GlCodeCombination.id == BudgetLine.code_combination_id)
            .where(BudgetLine.version_id == vid, GlCodeCombination.account_type == "E")
        ).all():
            outflow[(line.fund_value, line.period_num)] += Decimal(line.amount)

    running = {f: cash.get(f, Decimal("0")) for f in funds}
    rows: list[dict] = []
    y, m = start.year, start.month
    for _ in range(months):
        for f in funds:
            inf = inflow.get((f, y, m), Decimal("0"))
            outf = outflow.get((f, m), Decimal("0"))
            opening = running[f]
            ending = opening + inf - outf
            running[f] = ending
            rows.append({"period": f"{_MONTHS[m - 1]}-{y}", "fund": f,
                         "opening": opening.quantize(Decimal("0.01")),
                         "inflow": inf.quantize(Decimal("0.01")),
                         "outflow": outf.quantize(Decimal("0.01")),
                         "ending": ending.quantize(Decimal("0.01"))})
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return rows


def build_cash_flow_forecast_workbook(db, tenant_id, start: date, months, tenant_name) -> bytes:
    rows = cash_flow_forecast(db, tenant_id, start, months)
    wb = Workbook(); ws = wb.active; ws.title = "Cash Flow Forecast"
    ws["A1"] = f"{tenant_name} — Cash Flow Forecast ({months} months from {start.isoformat()})"
    ws["A1"].font = Font(size=14, bold=True)
    _header_row(ws, 3, ["Period", "Fund", "Opening", "Inflows", "Outflows", "Ending"])
    r = 4
    for row in rows:
        ws.cell(row=r, column=1, value=row["period"])
        ws.cell(row=r, column=2, value=row["fund"])
        ws.cell(row=r, column=3, value=float(row["opening"]))
        ws.cell(row=r, column=4, value=float(row["inflow"]))
        ws.cell(row=r, column=5, value=float(row["outflow"]))
        ws.cell(row=r, column=6, value=float(row["ending"]))
        r += 1
    for col, w in {"A": 12, "B": 8, "C": 14, "D": 14, "E": 14, "F": 14}.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def exec_dashboard(db: Session, tenant_id) -> dict:
    from app.models.collections import DelinquencyCase, Lien, PaymentPlan
    cash = _cash_by_fund(db, tenant_id)
    ar = _ar_open_by_fund(db, tenant_id)
    funds = sorted(set(cash) | set(ar))
    today = date.today()
    ag = aging(db, tenant_id, today)
    delinquent_total = sum((e["total"] for e in ag), Decimal("0"))
    open_cases = db.execute(select(func.count(DelinquencyCase.id)).where(
        DelinquencyCase.tenant_id == tenant_id, DelinquencyCase.stage != "RESOLVED")).scalar_one()
    active_plans = db.execute(select(func.count(PaymentPlan.id)).where(
        PaymentPlan.tenant_id == tenant_id, PaymentPlan.status == "ACTIVE")).scalar_one()
    filed_liens = db.execute(select(func.count(Lien.id)).where(
        Lien.tenant_id == tenant_id, Lien.status == "FILED")).scalar_one()
    return {
        "funds": [{"fund": f, "cash": str(cash.get(f, Decimal("0")).quantize(Decimal("0.01"))),
                   "ar_open": str(ar.get(f, Decimal("0")).quantize(Decimal("0.01")))} for f in funds],
        "cash_total": str(sum(cash.values(), Decimal("0")).quantize(Decimal("0.01"))),
        "ar_open_total": str(sum(ar.values(), Decimal("0")).quantize(Decimal("0.01"))),
        "delinquent_total": str(delinquent_total.quantize(Decimal("0.01"))),
        "open_cases": open_cases, "active_plans": active_plans, "filed_liens": filed_liens,
        "aging_by_fund": {f: {b: str(v.quantize(Decimal("0.01"))) for b, v in buckets.items()}
                          for f, buckets in aging_by_fund(db, tenant_id, today).items()},
    }


def build_delinquency_packet_pdf(db: Session, tenant_id, as_of: date, tenant_name, top_n=15) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title=f"{tenant_name} Delinquency Packet")
    styles = getSampleStyleSheet()
    el = [Paragraph(f"{tenant_name} — Delinquency Board Packet", styles["Title"]),
          Paragraph(f"As of {as_of.isoformat()}", styles["Normal"]), Spacer(1, 0.2 * inch)]

    # Aging by Fund.
    el.append(Paragraph("Aging by Fund", styles["Heading2"]))
    fund_rows = [["Fund", *BUCKETS]]
    for fund, b in sorted(aging_by_fund(db, tenant_id, as_of).items()):
        fund_rows.append([fund, *[f"${b[k]:,.0f}" for k in BUCKETS]])
    if len(fund_rows) == 1:
        fund_rows.append(["—"] + ["$0"] * len(BUCKETS))
    t = Table(fund_rows, hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                           ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    el += [t, Spacer(1, 0.2 * inch)]

    # Top delinquents.
    el.append(Paragraph(f"Top {top_n} Delinquent Accounts", styles["Heading2"]))
    ag = aging(db, tenant_id, as_of)[:top_n]
    rows = [["Account", "Homeowner", "90+", "Total"]]
    for e in ag:
        rows.append([e["account_number"], e["name"], f"${e['buckets']['90+']:,.0f}", f"${e['total']:,.2f}"])
    if len(rows) == 1:
        rows.append(["—", "No delinquencies", "$0", "$0"])
    t2 = Table(rows, hAlign="LEFT", colWidths=[1.2 * inch, 2.6 * inch, 1.0 * inch, 1.2 * inch])
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    el += [t2, Spacer(1, 0.2 * inch)]

    # Collection effectiveness (year to as_of).
    eff = collection_effectiveness(db, tenant_id, date(as_of.year, 1, 1), as_of)
    el.append(Paragraph("Collection Effectiveness (YTD)", styles["Heading2"]))
    el.append(Paragraph(
        f"Billed ${eff['billed']:,.2f} · Collected ${eff['collected']:,.2f} · "
        f"Rate {eff['rate_pct']}%", styles["Normal"]))

    doc.build(el)
    return buf.getvalue()
